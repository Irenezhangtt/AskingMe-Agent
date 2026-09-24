"""Upload preview -> reviewed policy -> date-aware price contract, no paid calls."""
import io
import json
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from PIL import Image
from pydantic import ValidationError

from api import main
from rag.policy_uploads import (PolicyEvidence, OrderFacts, UploadedPriceRequest,
                                extract_policy, calculate_uploaded_price, image_block)

TEXT = ('Valid from 2026-09-01 through 2026-09-30. Appliances only. CNY 50 off an original '
        'subtotal of at least CNY 300, once per order. Then apply a claimed and valid CNY 20 coupon. '
        'Members receive 10% off the remainder. Final subtotal excludes shipping and tax. '
        'Round once to two decimal places, half-up.')


def policy(**kwargs):
    return PolicyEvidence(**{'source_name': 'offer.md', 'title': 'September offer', 'content': TEXT,
        'date_scope': 'dated', 'effective_from': '2026-09-01', 'effective_to': '2026-09-30',
        'date_evidence': 'Valid from 2026-09-01 through 2026-09-30.', 'tags': ['Appliances'],
        'reviewed': True, **kwargs})


def order(**kwargs):
    return OrderFacts(**{'subtotal': '400.00', 'currency': 'CNY', 'purchase_date': '2026-09-22',
        'category': 'Appliances', 'member': True, 'coupon_claimed': True, 'coupon_valid': True, **kwargs})


def llm(result):
    response = SimpleNamespace(content=[SimpleNamespace(type='text', text=json.dumps(result))])
    return SimpleNamespace(model='test-vision', client=SimpleNamespace(messages=SimpleNamespace(create=AsyncMock(return_value=response))))


def extraction():
    return policy(reviewed=False).model_dump(mode='json', exclude={'source_name', 'reviewed'})


def plan():
    return {'type': 'calculation', 'expression': '(400 - 50 - 20) * 0.9', 'currency': 'CNY', 'source_ids': [1], 'rule_checks': [{'condition': 'Stacking order', 'outcome': 'applies', 'reason': 'Threshold then coupon then member discount.', 'source_ids': [1]}]}


class PolicyUploadTests(unittest.IsolatedAsyncioTestCase):
    async def test_text_extraction_preserves_complete_source_instead_of_model_summary(self):
        data = extraction()
        data['content'] = 'Misleading summary without exclusions'
        result = await extract_policy(llm(data), 'offer.md', text=TEXT)
        self.assertEqual(result.content, TEXT)
        self.assertFalse(result.reviewed)
        self.assertEqual(result.date_scope, 'dated')

    async def test_unsubstantiated_dates_require_review(self):
        data = extraction()
        data['date_evidence'] = 'Invented date quote'
        result = await extract_policy(llm(data), 'offer.md', text=TEXT)
        self.assertEqual(result.date_scope, 'unknown')
        self.assertIsNone(result.effective_from)
        self.assertTrue(result.warnings)

    async def test_image_validated_and_sent_to_vision_with_untrusted_data_guard(self):
        buffer = io.BytesIO()
        Image.new('RGB', (20, 20), 'white').save(buffer, 'PNG')
        model = llm(extraction())
        result = await extract_policy(model, 'offer.png', image=image_block(buffer.getvalue()))
        sent = model.client.messages.create.call_args.kwargs
        self.assertEqual(sent['messages'][0]['content'][0]['source']['media_type'], 'image/png')
        self.assertIn('never instructions', sent['system'])
        self.assertEqual(result.content, TEXT)
        with self.assertRaises(ValueError):
            image_block(b'not an image')

    async def test_date_boundaries_are_inclusive_and_arithmetic_is_computed(self):
        for day in ('2026-09-01', '2026-09-30'):
            result = await calculate_uploaded_price(llm(plan()), UploadedPriceRequest(policies=[policy()], order=order(purchase_date=day)))
            self.assertEqual(result['calculation']['final_price'], '297.00')
            self.assertEqual(len(result['calculation']['steps']), 3)

    async def test_outside_dates_excluded_without_model_or_invented_price(self):
        for day in ('2026-08-31', '2026-10-01'):
            model = llm(plan())
            result = await calculate_uploaded_price(model, UploadedPriceRequest(policies=[policy()], order=order(purchase_date=day)))
            self.assertEqual(result['type'], 'clarification')
            self.assertEqual(len(result['excluded']), 1)
            model.client.messages.create.assert_not_awaited()

    async def test_unreviewed_or_unknown_validity_blocks_calculation(self):
        for evidence in (policy(reviewed=False), policy(date_scope='unknown', effective_from=None, effective_to=None)):
            model = llm(plan())
            result = await calculate_uploaded_price(model, UploadedPriceRequest(policies=[evidence], order=order()))
            self.assertEqual(result['type'], 'clarification')
            model.client.messages.create.assert_not_awaited()

    async def test_all_current_policies_and_unknown_facts_reach_model(self):
        model = llm({'type': 'clarification', 'answer': 'Are you a member, and which promotion applies?'})
        result = await calculate_uploaded_price(model, UploadedPriceRequest(
            policies=[policy(), policy(source_name='conflicting.md', content='Cannot combine with any other offer.', tags=['Different label']),
                      policy(source_name='old.md', effective_from='2025-09-01', effective_to='2025-09-30')],
            order=order(member=None)))
        payload = json.loads(model.client.messages.create.call_args.kwargs['messages'][0]['content'])
        self.assertEqual(len(payload['Rule documents']), 2)
        self.assertIsNone(payload['Order']['member'])
        self.assertEqual(result['excluded'][0]['source_name'], 'old.md')
        self.assertNotIn('calculation', result)

    async def test_currency_and_invalid_plans_fail_closed(self):
        for bad in ({**plan(), 'currency': 'USD'}, {**plan(), 'expression': '1/0'}, {**plan(), 'source_ids': [99]}, {'type': 'explanation', 'answer': 'Pay 1 CNY'}, []):
            result = await calculate_uploaded_price(llm(bad), UploadedPriceRequest(policies=[policy()], order=order()))
            self.assertEqual(result['type'], 'clarification')
            self.assertNotIn('calculation', result)

    def test_bad_dates_and_amounts_rejected(self):
        with self.assertRaises(ValidationError):
            policy(effective_from='2026-10-01')
        for value in ('-1', 'NaN', '1.001', '1e2'):
            with self.assertRaises(ValidationError):
                order(subtotal=value)


class PolicyUploadAPITests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(main.app)
        main.app.dependency_overrides[main.require_admin_key] = lambda: None
        self.model = llm(extraction())
        self.patcher = patch.object(main, '_rag', SimpleNamespace(llm=self.model))
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        main.app.dependency_overrides.clear()

    def test_upload_review_calculate_without_shared_knowledge(self):
        with patch.object(main, '_tool_manager', None):
            upload = self.client.post('/pricing/extract', files=[('files', ('offer.md', TEXT.encode(), 'text/markdown'))])
            self.assertEqual(upload.status_code, 200, upload.text)
            evidence = upload.json()['policies']
            body = {'policies': evidence, 'order': order().model_dump(mode='json')}
            self.assertEqual(self.client.post('/pricing/calculate', json=body).json()['type'], 'clarification')
            evidence[0]['reviewed'] = True
            self.model.client.messages.create.return_value.content[0].text = json.dumps(plan())
            response = self.client.post('/pricing/calculate', json=body)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['calculation']['final_price'], '297.00')

    def test_authentication_required_for_both_endpoints(self):
        main.app.dependency_overrides.clear()
        with patch.dict('os.environ', {'ADMIN_API_KEY': 'test-secret'}):
            self.assertEqual(self.client.post('/pricing/extract', files={'files': ('offer.md', b'policy')}).status_code, 401)
            self.assertEqual(self.client.post('/pricing/calculate', json={'policies': [policy().model_dump(mode='json')], 'order': order().model_dump(mode='json')}).status_code, 401)

    def test_corrupt_oversize_and_long_files_make_no_model_calls(self):
        for filename, content, status in [('offer.png', b'not an image', 400), ('offer.md', b'x' * (5 * 1024 * 1024 + 1), 413), ('offer.md', b'x' * 16001, 400), ('offer.txt', b'\xff', 400)]:
            response = self.client.post('/pricing/extract', files={'files': (filename, content)})
            self.assertEqual(response.status_code, status, response.text)
        self.model.client.messages.create.assert_not_awaited()

    def test_upstream_failure_does_not_leak_provider_details(self):
        self.model.client.messages.create.side_effect = RuntimeError('private provider internals')
        response = self.client.post('/pricing/extract', files={'files': ('offer.md', TEXT.encode())})
        self.assertEqual(response.status_code, 502)
        self.assertNotIn('private provider internals', response.text)

    def test_docx_keeps_table_in_document_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Before discount')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'CNY 20 coupon'
        document.add_paragraph('After discount')
        buffer = io.BytesIO()
        document.save(buffer)
        content = main._extract_uploaded_documents('offer.docx', buffer.getvalue())[0]['content']
        self.assertLess(content.index('Before'), content.index('CNY'))
        self.assertLess(content.index('CNY'), content.index('After'))

    def test_blank_pdf_page_is_not_silently_dropped(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        response = self.client.post('/pricing/extract', files={'files': ('offer.pdf', buffer.getvalue())})
        self.assertEqual(response.status_code, 400)
        self.assertIn('screenshots', response.json()['detail'])
        self.model.client.messages.create.assert_not_awaited()


if __name__ == '__main__':
    unittest.main()
