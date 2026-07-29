"""FastAPI entry point for the AskingMe Agent enterprise policy Q&A system.

All core components are initialized during lifespan and configured through
environment variables.
"""
import asyncio
import hashlib
import io
import json
import logging
import os
import pathlib
import re
import secrets
import sys
import time
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from typing import Any, Awaitable, Callable, Dict, List, Optional


_ROOT = str(pathlib.Path(__file__).parent.parent.resolve())
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import uvicorn
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, Form, Header, HTTPException, Request, Response, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from prometheus_client import CONTENT_TYPE_LATEST, generate_latest
from pydantic import BaseModel, Field
from core.conversation_rules import is_conversation_memory_query

load_dotenv()

logging.basicConfig(
    level=getattr(logging, os.getenv("LOG_LEVEL", "INFO")),
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

BANNER = r"""
    ʕ•ᴥ•ʔ  ʕ•ᴥ•ʔ  ʕ•ᴥ•ʔ
   ╔══════════════════════╗
    ║  AskingMe Agent v1   ║
    ║ Enterprise Policy Assistant ║
   ╚══════════════════════╝
    ʕ•ᴥ•ʔ  ʕ•ᴥ•ʔ  ʕ•ᴥ•ʔ
"""

# ── Global components initialized during lifespan ─────────────────────────────
_orchestrator = None
_memory       = None
_tool_manager = None
_monitor      = None
_evaluator    = None
_skill_manager = None

def _anthropic_cfg() -> Dict[str, Any]:
    key = os.getenv("ANTHROPIC_API_KEY", "")
    if not key:
        raise RuntimeError("ANTHROPIC_API_KEY is not configured")
    cfg: Dict[str, Any] = {
        "api_key":  key,
        "model":    os.getenv(
            "ANSWER_MODEL",
            os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-6"),
        ).strip(),
        "router_model": os.getenv(
            "ROUTER_MODEL",
            os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-6"),
        ).strip(),
        "retrieval_model": os.getenv(
            "RETRIEVAL_MODEL",
            os.getenv("ROUTER_MODEL", os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-6")),
        ).strip(),
    }
    base_url = os.getenv("ANTHROPIC_BASE_URL", "").strip()
    if base_url:
        cfg["base_url"] = base_url
    return cfg


@asynccontextmanager
async def lifespan(app: FastAPI):
    global _orchestrator, _memory, _tool_manager, _monitor, _evaluator, _skill_manager

    print(BANNER, flush=True)

    from agents.agent_orchestrator import AgentOrchestrator, Request
    from core.intent_recognizer import IntentRecognizer
    from evaluation.evaluator import EndToEndEvaluator
    from mcp.knowledge_base import KnowledgeBase
    from mcp.tool_manager import MCPToolManager, Tool
    from memory.conversation_memory import MemoryManager
    from monitor.performance_monitor import PerformanceMonitor
    from core.skill_loader import SkillManager

    cfg = _anthropic_cfg()
    logger.info(
        "Models: answer=%s router=%s retrieval=%s base_url=%s",
        cfg["model"],
        cfg["router_model"],
        cfg["retrieval_model"],
        cfg.get("base_url", "(official)"),
    )

    # Intent recognizer; the orchestrator also creates one, while this instance is exposed to the evaluator.
    recognizer = IntentRecognizer(
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["router_model"],
    )

    # Skills: load business capabilities at startup and inject them into agent LLM calls.
    skills_dir = os.getenv("ASKINGME_SKILLS_DIR", str(pathlib.Path(_ROOT) / "skills"))
    _skill_manager = SkillManager(
        root_dir=skills_dir,
        max_prompt_chars=int(os.getenv("ASKINGME_SKILLS_MAX_PROMPT_CHARS", "5000")),
    )
    _skill_manager.load()

    # Agent orchestrator
    _orchestrator = AgentOrchestrator(
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["model"],
        router_model=cfg["router_model"],
        skill_manager=_skill_manager,
    )

    # Memory manager: Redis working memory plus ChromaDB episodic memory and user profiles
    _memory = MemoryManager(
        redis_url=os.getenv("REDIS_URL", "redis://redis:6379/0"),
        chroma_host=os.getenv("CHROMA_HOST", "chromadb"),
        chroma_port=int(os.getenv("CHROMA_PORT", "8000")),
        chroma_path=os.getenv("CHROMA_PERSIST_DIRECTORY", "/app/data/chroma"),
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["retrieval_model"],
    )

    # MCP tool manager and ChromaDB-backed RAG knowledge base
    _tool_manager = MCPToolManager(
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["model"],
    )
    kb = KnowledgeBase(
        chroma_host=os.getenv("CHROMA_HOST", "chromadb"),
        chroma_port=int(os.getenv("CHROMA_PORT", "8000")),
        chroma_path=os.getenv("CHROMA_PERSIST_DIRECTORY", "/app/data/chroma"),
    )
    logger.info(f"Knowledge base loaded: {kb.doc_count} document chunks")
    if os.getenv("RAG_WARMUP_ENABLED", "true").lower() == "true":
        try:
            await asyncio.to_thread(kb.search, "company policy approval process", 1)
            logger.info("Knowledge-base embedding path warmed up")
        except Exception as ex:
            logger.warning("Knowledge-base warmup failed: %s", ex)

    def knowledge_fallback(params: Dict[str, Any], context: Optional[Dict[str, Any]], error: str):
        query = params.get("query", "")
        return [{
            "title": "Knowledge-base fallback result",
            "content": f"The policy knowledge base is temporarily unavailable and could not search for \"{query}\". Try again later or confirm with the policy owner.",
            "score": 0.0,
            "fallback": True,
            "error": error,
        }]

    _tool_manager.register(Tool(
        name="knowledge_search",
        description="Search the policy knowledge base with ChromaDB vector retrieval",
        handler=kb.search_handler,
        schema={
            "type": "object",
            "properties": {
                "query": {"type": "string"},
                "top_k": {"type": "integer"},
            },
            "required": ["query"],
        },
        cache_ttl=300.0,
        supports_rerank=True,
        fallback=knowledge_fallback,
    ))

    # Performance monitoring with optional Prometheus startup
    prom_port = int(os.getenv("PROMETHEUS_PORT", "0")) or None
    _monitor = PerformanceMonitor(
        orchestrator=_orchestrator,
        tool_manager=_tool_manager,
        interval_s=float(os.getenv("MONITOR_INTERVAL", "10")),
        webhook_url=os.getenv("ALERT_WEBHOOK_URL") or None,
        prometheus_port=prom_port,
    )
    await _monitor.start()

    # Evaluator
    _evaluator = EndToEndEvaluator(
        orchestrator=_orchestrator,
        recognizer=recognizer,
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["model"],
        baseline_path=os.getenv("EVAL_BASELINE_PATH", "/app/data/eval/baseline.json"),
    )

    logger.info("AskingMe Agent is ready")
    yield

    await _monitor.stop()
    logger.info("AskingMe Agent has shut down")


# ── FastAPI ───────────────────────────────────────────────────────────────────
app = FastAPI(
    title="AskingMe Agent Enterprise Policy Assistant",
    version="2.0.0",
    lifespan=lifespan,
    docs_url="/docs",
)


def _csv_env(name: str, default: str) -> List[str]:
    """Read a comma-separated environment variable and ignore empty values."""
    return [item.strip() for item in os.getenv(name, default).split(",") if item.strip()]


app.add_middleware(
    CORSMiddleware,
    allow_origins=_csv_env(
        "CORS_ORIGINS",
        "http://localhost,http://127.0.0.1,http://localhost:5173,http://127.0.0.1:5173",
    ),
    allow_credentials=os.getenv("CORS_ALLOW_CREDENTIALS", "false").lower() == "true",
    allow_methods=_csv_env("CORS_ALLOW_METHODS", "GET,POST,OPTIONS"),
    allow_headers=_csv_env("CORS_ALLOW_HEADERS", "Content-Type,Authorization"),
)


# ── Request and response models ───────────────────────────────────────────────
class ChatRequest(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    user_id: str = Field(default="anonymous", min_length=1, max_length=128)
    conv_id: Optional[str] = Field(default=None, max_length=128)


class ChatResponse(BaseModel):
    conv_id:     str
    response:    str
    intent:      str
    agent_type:  str
    escalated:   bool
    latency_ms:  float
    knowledge_used: bool = False
    cache_hit: bool = False
    timings: Dict[str, float] = Field(default_factory=dict)


async def require_admin_key(
    x_admin_key: Optional[str] = Header(default=None, alias="X-Admin-Key"),
) -> None:
    """Protect administrative endpoints using a server-side environment key."""
    configured_key = os.getenv("ADMIN_API_KEY", "").strip()
    if not configured_key:
        logger.error("ADMIN_API_KEY is not configured; rejecting administrative request")
        raise HTTPException(503, "The administrative API is not configured")
    if not x_admin_key or not secrets.compare_digest(x_admin_key, configured_key):
        raise HTTPException(
            status_code=401,
            detail="Invalid administrator key",
            headers={"WWW-Authenticate": "ApiKey"},
        )


def _positive_int_env(name: str, default: int) -> int:
    try:
        return max(1, int(os.getenv(name, str(default))))
    except ValueError:
        logger.warning("%s is invalid; using default value %s", name, default)
        return default


def _demo_quota_context(request: Request) -> Dict[str, Any]:
    ip_limit = _positive_int_env("DEMO_DAILY_IP_LIMIT", 5)
    global_limit = _positive_int_env("DEMO_DAILY_GLOBAL_LIMIT", 50)
    now = datetime.now(timezone.utc)
    tomorrow = (now + timedelta(days=1)).replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )
    client_ip = request.headers.get("x-real-ip") or (
        request.client.host if request.client else "unknown"
    )
    ip_hash = hashlib.sha256(client_ip.encode("utf-8")).hexdigest()[:20]
    day = now.strftime("%Y-%m-%d")
    return {
        "global_limit": global_limit,
        "ip_limit": ip_limit,
        "global_key": f"demo:quota:{day}:global",
        "ip_key": f"demo:quota:{day}:ip:{ip_hash}",
        "reset_at": tomorrow,
        "ttl_seconds": max(1, int((tomorrow - now).total_seconds())),
    }


async def enforce_demo_quota(request: Request, response: Response) -> None:
    """Use atomic Redis counters to enforce global and per-IP demo quotas."""
    if os.getenv("DEMO_QUOTA_ENABLED", "true").lower() != "true":
        return
    if _memory is None:
        raise HTTPException(503, "Service is not ready")

    quota = _demo_quota_context(request)

    script = """
    local global_count = redis.call('INCR', KEYS[1])
    if global_count == 1 then redis.call('EXPIRE', KEYS[1], ARGV[3]) end
    if global_count > tonumber(ARGV[1]) then
        redis.call('DECR', KEYS[1])
        return {0, global_count - 1, -1}
    end

    local ip_count = redis.call('INCR', KEYS[2])
    if ip_count == 1 then redis.call('EXPIRE', KEYS[2], ARGV[3]) end
    if ip_count > tonumber(ARGV[2]) then
        redis.call('DECR', KEYS[1])
        return {-1, global_count - 1, ip_count}
    end
    return {1, global_count, ip_count}
    """

    try:
        allowed, global_count, ip_count = _memory._redis.eval(
            script,
            2,
            quota["global_key"],
            quota["ip_key"],
            quota["global_limit"],
            quota["ip_limit"],
            quota["ttl_seconds"],
        )
    except Exception as ex:
        logger.error("Demo quota check failed: %s", ex)
        raise HTTPException(503, "The demo quota service is temporarily unavailable") from ex

    quota_headers = {
        "X-RateLimit-Limit": str(quota["ip_limit"]),
        "X-RateLimit-Remaining": str(
            max(0, quota["ip_limit"] - max(0, ip_count))
        ),
        "X-RateLimit-Reset": str(int(quota["reset_at"].timestamp())),
    }
    response.headers.update(quota_headers)

    if allowed == 0:
        raise HTTPException(
            429,
            "The public demo's global quota has been reached for today. Please try again tomorrow.",
            headers={**quota_headers, "Retry-After": str(quota["ttl_seconds"])},
        )
    if allowed == -1:
        raise HTTPException(
            429,
            "You have reached your demo quota for today. Please try again tomorrow.",
            headers={
                **quota_headers,
                "X-RateLimit-Remaining": "0",
                "Retry-After": str(quota["ttl_seconds"]),
            },
        )


# ── Routes ────────────────────────────────────────────────────────────────────
@app.get("/health")
async def health():
    if _orchestrator is None or _memory is None:
        raise HTTPException(503, "Service is not ready")
    components = _memory.health_status()
    components["agents"] = True
    healthy = all(components.values())
    return JSONResponse(
        status_code=200 if healthy else 503,
        content={
            "status": "ok" if healthy else "degraded",
            "components": components,
            "agents": _orchestrator.get_stats(),
        },
    )


@app.get("/demo/quota")
async def demo_quota(request: Request):
    """Return the current visitor's demo quota without consuming a call."""
    enabled = os.getenv("DEMO_QUOTA_ENABLED", "true").lower() == "true"
    if not enabled:
        return {"enabled": False}
    if _memory is None:
        raise HTTPException(503, "Service is not ready")

    quota = _demo_quota_context(request)
    try:
        ip_count = int(_memory._redis.get(quota["ip_key"]) or 0)
        global_count = int(_memory._redis.get(quota["global_key"]) or 0)
    except Exception as ex:
        logger.error("Failed to read demo quota: %s", ex)
        raise HTTPException(503, "The demo quota service is temporarily unavailable") from ex

    return {
        "enabled": True,
        "limit": quota["ip_limit"],
        "remaining": max(0, quota["ip_limit"] - ip_count),
        "available": global_count < quota["global_limit"],
        "reset_at": quota["reset_at"].isoformat(),
    }


@app.get("/skills", tags=["Skills"], dependencies=[Depends(require_admin_key)])
async def skills_summary():
    """Return loaded skills for reload verification and parser diagnostics."""
    if _skill_manager is None:
        raise HTTPException(503, "Skills are not initialized")
    return _skill_manager.summary()


@app.post("/skills/reload", tags=["Skills"], dependencies=[Depends(require_admin_key)])
async def reload_skills():
    """Rescan the skill directory at runtime without restarting the service."""
    if _skill_manager is None:
        raise HTTPException(503, "Skills are not initialized")
    _skill_manager.reload()
    if _orchestrator is not None:
        _orchestrator.set_skill_manager(_skill_manager)
    return _skill_manager.summary()


@app.post(
    "/chat",
    response_model=ChatResponse,
    dependencies=[Depends(enforce_demo_quota)],
)
async def chat(req: ChatRequest):
    """Run the optimized chat pipeline and return one JSON response."""
    return await _process_chat(req)


@app.post(
    "/chat/stream",
    dependencies=[Depends(enforce_demo_quota)],
)
async def chat_stream(req: ChatRequest):
    """Stream processing phases and deliver the final answer incrementally."""
    queue: asyncio.Queue[Optional[Dict[str, Any]]] = asyncio.Queue()

    async def progress(event: Dict[str, Any]) -> None:
        await queue.put(event)

    async def run_pipeline() -> None:
        try:
            result = await _process_chat(req, progress=progress)
            payload = result.model_dump()
            response_text = payload.pop("response")
            await queue.put({"type": "meta", "data": payload})
            # Cached answers have no upstream model stream, so chunk them here.
            if result.cache_hit:
                chunk_size = _positive_int_env("STREAM_CHUNK_CHARS", 48)
                for start in range(0, len(response_text), chunk_size):
                    await queue.put({
                        "type": "answer",
                        "delta": response_text[start:start + chunk_size],
                    })
                    await asyncio.sleep(0)
            await queue.put({"type": "done"})
        except HTTPException as ex:
            await queue.put({"type": "error", "status": ex.status_code, "detail": ex.detail})
        except Exception:
            logger.exception("Streaming chat pipeline failed")
            await queue.put({"type": "error", "status": 500, "detail": "The AI service is temporarily unavailable"})
        finally:
            await queue.put(None)

    async def events():
        task = asyncio.create_task(run_pipeline())
        try:
            while True:
                event = await queue.get()
                if event is None:
                    break
                yield json.dumps(event, ensure_ascii=False) + "\n"
        finally:
            if not task.done():
                task.cancel()

    return StreamingResponse(
        events(),
        media_type="application/x-ndjson",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "X-Accel-Buffering": "no",
        },
    )


async def _process_chat(
    req: ChatRequest,
    progress: Optional[Callable[[Dict[str, Any]], Awaitable[None]]] = None,
) -> ChatResponse:
    """Run cache, memory, adaptive RAG, orchestration, and persistence."""
    if _orchestrator is None or _memory is None:
        raise HTTPException(503, "Service is not ready")

    from agents.agent_orchestrator import Request as OrcReq
    from memory.conversation_memory import MsgRole

    total_t0 = time.monotonic()
    conv_id = req.conv_id or str(uuid.uuid4())
    timings: Dict[str, float] = {}

    async def emit(phase: str, detail: str) -> None:
        if progress is not None:
            await progress({"type": "phase", "phase": phase, "detail": detail})

    await emit("cache", "Checking the policy answer cache")
    cached = _get_cached_answer(req.message) if _is_answer_cacheable(req) else None
    if cached is not None:
        await emit("cache_hit", "A verified cached answer was found")
        await _memory.add_message(req.user_id, conv_id, MsgRole.USER, req.message)
        await _memory.add_message(req.user_id, conv_id, MsgRole.ASSISTANT, cached["response"])
        total_ms = (time.monotonic() - total_t0) * 1000
        return ChatResponse(
            conv_id=conv_id,
            response=cached["response"],
            intent=cached.get("intent", "query"),
            agent_type=cached.get("agent_type", "general"),
            escalated=bool(cached.get("escalated", False)),
            latency_ms=round(total_ms, 1),
            knowledge_used=bool(cached.get("knowledge_used", True)),
            cache_hit=True,
            timings={"cache_ms": round(total_ms, 1), "total_ms": round(total_ms, 1)},
        )

    await emit("context", "Reading conversation memory and retrieving policy context")

    async def load_memory():
        started = time.monotonic()
        value = await _memory.get_context(req.user_id, conv_id, query=req.message)
        return value, (time.monotonic() - started) * 1000

    async def load_knowledge():
        started = time.monotonic()
        value = await _build_knowledge_context(req.message)
        return value, (time.monotonic() - started) * 1000

    # Memory and policy retrieval are independent, so run them concurrently.
    (mem_ctx, memory_ms), ((knowledge_text, knowledge_used), retrieval_ms) = await asyncio.gather(
        load_memory(),
        load_knowledge(),
    )
    timings["memory_ms"] = round(memory_ms, 1)
    timings["retrieval_ms"] = round(retrieval_ms, 1)

    history = [
        {"role": m.role.value, "content": m.content}
        for m in mem_ctx.recent_messages[-5:]
    ] if mem_ctx.recent_messages else None

    context_parts = [mem_ctx.to_prompt_text()]
    if knowledge_text:
        context_parts.append(knowledge_text)
    full_context = "\n\n".join(part for part in context_parts if part)

    orch_req = OrcReq(
        message=req.message,
        user_id=req.user_id,
        conv_id=conv_id,
        context=full_context,
        history=history,
        token_callback=(
            (lambda delta: progress({"type": "answer", "delta": delta}))
            if progress is not None
            else None
        ),
    )

    await emit("answer", "Routing to the policy specialist and drafting the answer")
    result = await _orchestrator.run(orch_req)
    timings.update(result.timings)

    await _memory.add_message(req.user_id, conv_id, MsgRole.USER, req.message)
    await _memory.add_message(req.user_id, conv_id, MsgRole.ASSISTANT, result.response)

    asyncio.create_task(_memory.update_profile(req.user_id, conv_id))
    total_ms = (time.monotonic() - total_t0) * 1000
    timings["total_ms"] = round(total_ms, 1)

    response = ChatResponse(
        conv_id=conv_id,
        response=result.response,
        intent=result.intent.value if result.intent else "other",
        agent_type=result.agent_type.value,
        escalated=result.escalated,
        latency_ms=round(total_ms, 1),
        knowledge_used=knowledge_used,
        timings=timings,
    )
    if _is_answer_cacheable(req) and knowledge_used and not result.escalated:
        _set_cached_answer(req.message, response)
    return response


def _is_answer_cacheable(req: ChatRequest) -> bool:
    """Cache only standalone policy questions whose answer is context independent."""
    return (
        os.getenv("ANSWER_CACHE_ENABLED", "true").lower() == "true"
        and req.conv_id is None
        and _should_use_knowledge(req.message)
        and not is_conversation_memory_query(req.message)
    )


def _answer_cache_key(message: str) -> str:
    normalized = " ".join((message or "").lower().split())
    policy_version = os.getenv("POLICY_KB_VERSION", "demo-v2")
    digest = hashlib.sha256(f"{policy_version}:{normalized}".encode("utf-8")).hexdigest()
    return f"askingme:answer:{digest}"


def _get_cached_answer(message: str) -> Optional[Dict[str, Any]]:
    if _memory is None:
        return None
    try:
        raw = _memory._redis.get(_answer_cache_key(message))
        return json.loads(raw) if raw else None
    except Exception as ex:
        logger.warning("Answer-cache read failed: %s", ex)
        return None


def _set_cached_answer(message: str, response: ChatResponse) -> None:
    if _memory is None:
        return
    try:
        ttl = _positive_int_env("ANSWER_CACHE_TTL_SECONDS", 3600)
        payload = response.model_dump(exclude={"conv_id", "latency_ms", "timings", "cache_hit"})
        _memory._redis.setex(
            _answer_cache_key(message),
            ttl,
            json.dumps(payload, ensure_ascii=False),
        )
    except Exception as ex:
        logger.warning("Answer-cache write failed: %s", ex)


def _invalidate_answer_cache() -> None:
    """Remove policy answers cached before a knowledge-base mutation."""
    if _memory is None:
        return
    try:
        keys = list(_memory._redis.scan_iter(match="askingme:answer:*", count=200))
        if keys:
            _memory._redis.delete(*keys)
        logger.info("Invalidated %s cached policy answers", len(keys))
    except Exception as ex:
        logger.warning("Answer-cache invalidation failed: %s", ex)


async def _build_knowledge_context(message: str, top_k: int = 3) -> tuple[str, bool]:
    """Build RAG context with direct retrieval and adaptive expansion."""
    if _tool_manager is None:
        return "", False
    if not _should_use_knowledge(message):
        return "", False
    try:
        result = await _tool_manager.search_optimized("knowledge_search", message, top_k=top_k)
        if not result.success or not isinstance(result.data, list) or not result.data:
            return "", False

        parts = ["[Knowledge-base retrieval results]"]
        used = False
        for i, item in enumerate(result.data[:top_k], start=1):
            if not isinstance(item, dict):
                continue
            title = str(item.get("title", "Untitled document"))
            content = str(item.get("content", "")).strip()
            score = item.get("score", "")
            if not content:
                continue
            used = True
            parts.append(f"{i}. Title: {title}\n   Relevance: {score}\n   Content: {content[:600]}")

        if not used:
            return "", False
        parts.append(
            "Base the answer on the policy provisions above and explain scope, process, and required documents. "
            "If information is insufficient or versions conflict, recommend the appropriate policy owner and do not invent an answer."
        )
        return "\n".join(parts), True
    except Exception as ex:
        logger.warning(f"Failed to build knowledge-base context: {ex}")
        return "", False


def _should_use_knowledge(message: str) -> bool:
    """Retrieve policy knowledge only for messages requiring business facts or rules."""
    msg = re.sub(r"\s+", "", (message or "").strip().lower())
    if not msg:
        return False

    # Strip common trailing punctuation so short greetings do not trigger retrieval.
    plain = msg.strip("，。！？,.!?~～")
    greetings = {
        "hi", "hello", "hey", "good morning", "good evening",
        "thanks", "thank you", "okay", "goodbye",
    }
    if plain in greetings:
        return False

    # Conversation-recall questions should use Redis or summaries instead of policy retrieval.
    if is_conversation_memory_query(plain):
        return False

    # Retrieve only for explicit business domains or rule-oriented expressions.
    knowledge_keywords = [
        "policy", "rule", "standard", "process", "workflow", "approval", "application", "documents", "effective",
        "version", "scope", "reimbursement", "travel", "hotel", "transportation", "meal allowance",
        "receipt", "expense", "procurement", "contract", "budget", "leave", "annual leave", "sick leave",
        "personal leave", "attendance", "timecard", "overtime", "compensatory leave", "benefits", "onboarding",
        "offboarding", "transfer", "access", "account", "data", "system", "role",
        "owner", "direct manager", "department head", "hr",
        "leave", "attendance", "access", "permission", "policy",
        "制度", "政策", "规定", "规则", "标准", "流程", "审批", "申请", "材料", "生效",
        "版本", "范围", "报销", "差旅", "出差", "住宿", "交通", "餐补", "发票", "费用",
        "采购", "合同", "预算", "请假", "年假", "病假", "事假", "考勤", "打卡", "加班",
        "调休", "福利", "入职", "离职", "转岗", "权限", "账号", "账户", "数据", "系统", "角色",
    ]
    return any(keyword in plain for keyword in knowledge_keywords)


@app.get("/monitor", dependencies=[Depends(require_admin_key)])
async def monitor_summary():
    """Return live agent, tool, alert, and optimization metrics."""
    if _monitor is None:
        raise HTTPException(503, "Service is not ready")
    return _monitor.summary()


@app.get("/metrics")
async def prometheus_metrics():
    """Expose Prometheus metrics."""
    return Response(generate_latest(), media_type=CONTENT_TYPE_LATEST)


@app.post("/search", dependencies=[Depends(require_admin_key)])
async def search(query: str, top_k: int = 5):
    """Demonstrate query rewriting, parallel retrieval, reranking, and top-K selection."""
    if _tool_manager is None:
        raise HTTPException(503, "Service is not ready")
    result = await _tool_manager.search_with_rewrite("knowledge_search", query, top_k=top_k)
    return {"query": query, "results": result.data, "reranked": result.reranked}


class DocInput(BaseModel):
    """Input model for one document."""
    title:   str
    content: str


class BatchDocInput(BaseModel):
    """Request model for batch document import."""
    documents: List[DocInput]


class EvalIntentInput(BaseModel):
    """Intent-recognition evaluation case."""
    message: str
    expected_intent: str
    context: Optional[Dict[str, Any]] = None


class EvalDialogInput(BaseModel):
    """Dialog-quality case supporting a single question or multiple turns."""
    question: Optional[str] = None
    turns: Optional[List[str]] = None
    user_id: Optional[str] = None
    conv_id: Optional[str] = None


class EvalRunInput(BaseModel):
    """Evaluation request that uses built-in cases when empty."""
    intent_cases: Optional[List[EvalIntentInput]] = None
    dialog_cases: Optional[List[EvalDialogInput]] = None


@app.post("/knowledge/add", tags=["Knowledge Base"], dependencies=[Depends(require_admin_key)])
async def add_knowledge(body: BatchDocInput):
    """Import, chunk, and embed documents in ChromaDB."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    kb = tool.handler.__self__
    count = kb.add_documents([{"title": d.title, "content": d.content} for d in body.documents])
    _invalidate_answer_cache()
    return {"message": f"Imported {count} document chunks", "added_chunks": count, "total_chunks": kb.doc_count}


def _extract_uploaded_documents(filename: str, content: bytes, title: str = "") -> List[Dict[str, str]]:
    """Extract text-oriented policy documents from supported upload formats."""
    suffix = pathlib.Path(filename).suffix.lower()
    default_title = title.strip() or pathlib.Path(filename).stem or "Untitled policy"

    if suffix == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        except Exception as ex:
            raise HTTPException(400, f"Failed to parse PDF: {ex}") from ex
        if not text.strip():
            raise HTTPException(400, "The PDF contains no extractable text; scanned PDFs require OCR")
        return [{"title": default_title, "content": text}]

    if suffix == ".docx":
        from docx import Document
        try:
            document = Document(io.BytesIO(content))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        except Exception as ex:
            raise HTTPException(400, f"Failed to parse DOCX: {ex}") from ex
        if not text.strip():
            raise HTTPException(400, "The DOCX file contains no extractable text")
        return [{"title": default_title, "content": text}]

    text = content.decode("utf-8", errors="ignore")
    if suffix == ".json":
        try:
            docs = json.loads(text)
            if not isinstance(docs, list):
                raise HTTPException(400, "JSON must be an array: [{title, content}, ...]")
            if not all(isinstance(doc, dict) and doc.get("title") and doc.get("content") for doc in docs):
                raise HTTPException(400, "Every JSON item requires non-empty title and content fields")
            return docs
        except json.JSONDecodeError as ex:
            raise HTTPException(400, f"Failed to parse JSON: {ex}") from ex

    if suffix not in {".txt", ".md"}:
        raise HTTPException(415, "Supported files are PDF, DOCX, TXT, MD, and JSON")
    if not text.strip():
        raise HTTPException(400, "The uploaded file is empty")
    return [{"title": default_title, "content": text}]


@app.post("/knowledge/upload", tags=["Knowledge Base"], dependencies=[Depends(require_admin_key)])
async def upload_knowledge(
    file: UploadFile = File(...),
    title: str = Form(""),
    version: str = Form("1.0"),
    approve: bool = Form(False),
    replaces_document_id: str = Form(""),
):
    """Parse and import a versioned policy file as a draft or approved document."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    kb = tool.handler.__self__

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "File exceeds the 10 MB limit")

    filename = file.filename or "unknown"
    docs = _extract_uploaded_documents(filename, content, title)
    from mcp.knowledge_base import DuplicateDocumentError
    imported = []
    try:
        for doc in docs:
            imported.append(kb.import_document(
                title=str(doc["title"]),
                content=str(doc["content"]),
                version=str(doc.get("version", version)),
                status="approved" if approve else "draft",
                source_name=filename,
                replaces_document_id=replaces_document_id,
            ))
    except DuplicateDocumentError as ex:
        raise HTTPException(
            409,
            {
                "message": str(ex),
                "duplicate": ex.document,
            },
        ) from ex
    except ValueError as ex:
        raise HTTPException(400, str(ex)) from ex

    if approve:
        _invalidate_answer_cache()
    return {
        "message": f"Imported file {filename}",
        "documents": imported,
        "added_chunks": sum(item["added_chunks"] for item in imported),
        "total_chunks": kb.doc_count,
    }


@app.get("/knowledge/stats", tags=["Knowledge Base"], dependencies=[Depends(require_admin_key)])
async def knowledge_stats():
    """Return knowledge-base statistics, including the total chunk count."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    kb = tool.handler.__self__
    documents = kb.list_documents()
    status_counts: Dict[str, int] = {}
    for document in documents:
        status = document["status"]
        status_counts[status] = status_counts.get(status, 0) + 1
    return {
        "total_chunks": kb.doc_count,
        "total_documents": len(documents),
        "status_counts": status_counts,
    }


@app.get("/knowledge/documents", tags=["Knowledge Base"], dependencies=[Depends(require_admin_key)])
async def list_knowledge_documents():
    """List document versions and lifecycle status."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    return {"documents": tool.handler.__self__.list_documents()}


@app.post(
    "/knowledge/documents/{document_id}/approve",
    tags=["Knowledge Base"],
    dependencies=[Depends(require_admin_key)],
)
async def approve_knowledge_document(document_id: str):
    """Approve a draft and archive older approved versions of the same policy."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    try:
        document = tool.handler.__self__.approve_document(document_id)
    except KeyError as ex:
        raise HTTPException(404, "Document not found") from ex
    _invalidate_answer_cache()
    return {"message": "Document approved", "document": document}


@app.delete(
    "/knowledge/documents/{document_id}",
    tags=["Knowledge Base"],
    dependencies=[Depends(require_admin_key)],
)
async def delete_knowledge_document(document_id: str):
    """Delete a document version and all of its vector chunks."""
    tool = _tool_manager._tools.get("knowledge_search") if _tool_manager else None
    if tool is None:
        raise HTTPException(503, "Knowledge base is not initialized")
    try:
        document = tool.handler.__self__.delete_document(document_id)
    except KeyError as ex:
        raise HTTPException(404, "Document not found") from ex
    if document.get("status") == "approved":
        _invalidate_answer_cache()
    return {"message": "Document deleted", "document": document}


@app.post("/eval/run", dependencies=[Depends(require_admin_key)])
async def run_eval(body: Optional[EvalRunInput] = None):
    """Run evaluation cases and return the report."""
    if _evaluator is None:
        raise HTTPException(503, "Service is not ready")
    from evaluation.evaluator import DEFAULT_DIALOG_CASES, DEFAULT_INTENT_CASES, IntentTestCase

    if body and body.intent_cases is not None:
        intent_cases = [
            IntentTestCase(
                message=c.message,
                expected_intent=c.expected_intent,
                context=c.context,
            )
            for c in body.intent_cases
        ]
    else:
        intent_cases = DEFAULT_INTENT_CASES

    if body and body.dialog_cases is not None:
        dialog_cases = [
            c.model_dump(exclude_none=True)
            for c in body.dialog_cases
        ]
    else:
        dialog_cases = DEFAULT_DIALOG_CASES

    report = await _evaluator.run(
        intent_cases=intent_cases,
        dialog_cases=dialog_cases,
    )
    return {
        "pass_rate":       report.pass_rate,
        "total":           report.total,
        "passed":          report.passed,
        "avg_scores":      report.avg_scores,
        "regressions":     report.regressions,
        "recommendations": report.recommendations,
        "results": [
            {
                "test_id": r.test_id,
                "passed": r.passed,
                "scores": r.scores,
                "detail": r.detail,
                "metadata": r.metadata,
            }
            for r in report.results
        ],
    }


# ── Interactive CLI ───────────────────────────────────────────────────────────
async def _cli():
    print(BANNER)
    print("AskingMe Agent CLI — type quit to exit\n")

    from agents.agent_orchestrator import AgentOrchestrator, Request
    from memory.conversation_memory import MemoryManager, MsgRole
    from core.skill_loader import SkillManager

    cfg = _anthropic_cfg()
    skill_manager = SkillManager(
        root_dir=os.getenv("ASKINGME_SKILLS_DIR", str(pathlib.Path(_ROOT) / "skills")),
        max_prompt_chars=int(os.getenv("ASKINGME_SKILLS_MAX_PROMPT_CHARS", "5000")),
    )
    skill_manager.load()
    orch = AgentOrchestrator(
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["model"],
        skill_manager=skill_manager,
    )
    mem  = MemoryManager(
        redis_url=os.getenv("REDIS_URL", "redis://localhost:6379/0"),
        chroma_host=os.getenv("CHROMA_HOST", "localhost"),
        chroma_port=int(os.getenv("CHROMA_PORT", "8000")),
        chroma_path=os.getenv("CHROMA_PERSIST_DIRECTORY", "/tmp/chroma"),
        api_key=cfg["api_key"],
        base_url=cfg.get("base_url"),
        model=cfg["model"],
    )

    user_id, conv_id = "cli_user", str(uuid.uuid4())

    while True:
        try:
            msg = input("You: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nGoodbye ʕ•ᴥ•ʔ")
            break
        if not msg or msg.lower() in ("quit", "exit"):
            print("Goodbye ʕ•ᴥ•ʔ")
            break

        ctx = await mem.get_context(user_id, conv_id, query=msg)
        history = [
            {"role": m.role.value, "content": m.content}
            for m in ctx.recent_messages[-5:]
        ] if ctx.recent_messages else None
        req = Request(message=msg, user_id=user_id, conv_id=conv_id, context=ctx.to_prompt_text(), history=history)
        result = await orch.run(req)

        await mem.add_message(user_id, conv_id, MsgRole.USER, msg)
        await mem.add_message(user_id, conv_id, MsgRole.ASSISTANT, result.response)

        print(f"\nAskingMe Agent [{result.agent_type.value}]: {result.response}\n")


if __name__ == "__main__":
    if "--cli" in sys.argv:
        asyncio.run(_cli())
    else:
        uvicorn.run(
            "api.main:app",
            host=os.getenv("API_HOST", "0.0.0.0"),
            port=int(os.getenv("API_PORT", "8000")),
            reload=os.getenv("APP_ENV") == "development",
        )
